#!/usr/bin/env python3
"""Seed/import FILME-V1.1.xlsm into the movie database.

Design goals:
- dry-run by default, no database connection required
- explicit --commit before any write happens
- cleaning/duplicate reports before writes
- async SQLAlchemy on commit
- no import of personal "Eigene Bewertung" / "Gesehen?" into UserMovie;
  those values are exported to a separate CSV report for later user assignment
- optional poster/person-image enrichment is OFF by default and requires
  --include-images

Expected application models (from the agreed V1 schema):
    app.models.Movie
    app.models.Person
    app.models.MoviePerson
    app.models.Genre
    app.models.MovieGenre
    app.models.Country
    app.models.MovieCountry
    app.models.ExternalRating
    app.models.MovieLink
and app.db.session.async_session_factory.
"""

from __future__ import annotations

import argparse
import asyncio
import csv
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import datetime, time, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

PROJECT_ROOT = Path(__file__).resolve().parents[1]
BACKEND_DIR = PROJECT_ROOT / "backend"
if BACKEND_DIR.exists() and str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

try:
    from python_calamine import CalamineWorkbook
except ImportError:
    from openpyxl import load_workbook

    class _OpenpyxlSheet:
        def __init__(self, worksheet):
            self.worksheet = worksheet

        def to_python(self, skip_empty_area: bool = False):
            return [list(row) for row in self.worksheet.iter_rows(values_only=True)]

    class CalamineWorkbook:  # compatible fallback used by this importer
        def __init__(self, workbook):
            self._workbook = workbook
            self.sheet_names = list(workbook.sheetnames)

        @classmethod
        def from_path(cls, path: str):
            workbook = load_workbook(
                path,
                read_only=True,
                data_only=True,
                keep_vba=True,
            )
            return cls(workbook)

        def get_sheet_by_name(self, name: str):
            return _OpenpyxlSheet(self._workbook[name])


MAIN_SHEET = "Filme"
POSTER_SHEET = "Fotos FP"
PERSON_IMAGE_SHEET = "Fotos R+S"

REQUIRED_HEADERS = {
    "Titel",
    "Originaltitel",
    "Genre",
    "Jahr",
    "Playtime",
    "Regisseur",
    "Hauptdarsteller 1",
    "Hauptdarsteller 2",
    "Hauptdarsteller 3",
    "Produktionsland",
    "IMDB (0-10)",
    "Tomatometer",
    "Popcornmeter",
    "Eigene Bewertung",
    "Gesehen?",
    "Kurzbeschreibung",
    "Anmerkung",
    "JustWatch",
    "ID",
}

NULL_STRINGS = {"", "-", "–", "—", "n/a", "na", "null", "none"}
GENRE_SEPARATOR = re.compile(r"\s*/\s*")
COUNTRY_SEPARATOR = re.compile(r"\s*,\s*")

# Values are deliberately only structural/common synonyms. Project-specific typo
# fixes belong in the overrides JSON, not hidden in code.
DEFAULT_SOURCE_NAMES = {
    "imdb": "imdb",
    "rotten_tomatoes": "rotten_tomatoes",
}


@dataclass(slots=True)
class WarningItem:
    row: int
    excel_id: str | None
    field: str
    value: Any
    code: str
    message: str
    severity: str = "warning"


@dataclass(slots=True)
class ParsedMovie:
    row: int
    excel_id: str | None
    title: str
    original_title: str | None
    release_year: int | None
    runtime_minutes: int | None
    overview: str | None
    editorial_note: str | None
    genres: list[str]
    countries: list[str]
    director: str | None
    cast: list[str]
    imdb: Decimal | None
    tomatometer: Decimal | None
    popcornmeter: Decimal | None
    justwatch_url: str | None
    own_rating_raw: str | None
    own_rating_0_100: int | None
    watched: bool
    poster_url: str | None = None

    @property
    def duplicate_key(self) -> tuple[str, int | None, str]:
        return (
            normalize_for_match(self.title),
            self.release_year,
            normalize_for_match(self.original_title or ""),
        )


@dataclass(slots=True)
class Overrides:
    genre_map: dict[str, str | list[str]]
    country_map: dict[str, str | list[str]]
    person_map: dict[str, str]
    title_map: dict[str, str]

    @classmethod
    def empty(cls) -> "Overrides":
        return cls({}, {}, {}, {})


def clean_optional_string(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    value = str(value).strip()
    if value.casefold() in NULL_STRINGS:
        return None
    return value


def normalize_for_match(value: str) -> str:
    value = unicodedata.normalize("NFKC", value).casefold().strip()
    value = re.sub(r"\s+", " ", value)
    return value


def parse_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return int(value)
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def parse_runtime_minutes(value: Any) -> int | None:
    """Convert Excel time/duration or a numeric Excel-day fraction to minutes."""
    if value is None or value == "":
        return None

    if isinstance(value, timedelta):
        return round(value.total_seconds() / 60)

    if isinstance(value, time):
        return value.hour * 60 + value.minute + round(value.second / 60)

    if isinstance(value, datetime):
        t = value.time()
        return t.hour * 60 + t.minute + round(t.second / 60)

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Excel stores time as a fraction of a day.
        if 0 <= float(value) < 10:
            return round(float(value) * 24 * 60)
        # Fallback in case a future sheet stores literal minutes.
        return round(float(value))

    text = str(value).strip()
    match = re.fullmatch(r"(?:(\d+):)?(\d{1,2}):(\d{2})", text)
    if match:
        hours = int(match.group(1) or 0)
        minutes = int(match.group(2))
        seconds = int(match.group(3))
        return hours * 60 + minutes + round(seconds / 60)

    try:
        numeric = float(text.replace(",", "."))
        return parse_runtime_minutes(numeric)
    except ValueError:
        return None


def parse_decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    text = clean_optional_string(value)
    if text is None:
        return None
    try:
        return Decimal(text.replace(",", "."))
    except Exception:
        return None


def parse_rt_percent(value: Any) -> Decimal | None:
    rating = parse_decimal(value)
    if rating is None:
        return None
    # Excel stores these as 0..1 fractions in the current workbook.
    if Decimal("0") <= rating <= Decimal("1"):
        return (rating * Decimal("100")).quantize(Decimal("0.001"))
    return rating.quantize(Decimal("0.001"))


def parse_watched(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    return str(value).strip().casefold() in {"1", "true", "yes", "ja", "x"}


def parse_star_rating(value: Any) -> tuple[str | None, int | None]:
    """Map legacy 0..5 star/half-star strings to 0..100.

    This is only exported as a pending personal-data report; it is not inserted
    into UserMovie by this script.
    """
    raw = clean_optional_string(value)
    if raw is None:
        return None, None

    full = raw.count("★")
    half = raw.count("⯪")
    if full == 0 and half == 0:
        return raw, None

    bricks = full + 0.5 * half
    if not 0 <= bricks <= 5:
        return raw, None
    return raw, int(round(bricks * 20))


def _expand_override(parts: list[str], mapping: dict[str, str | list[str]]) -> list[str]:
    result: list[str] = []
    for part in parts:
        mapped = mapping.get(part, part)
        if isinstance(mapped, list):
            result.extend(str(item).strip() for item in mapped if str(item).strip())
        else:
            value = str(mapped).strip()
            if value:
                result.append(value)
    return result


def split_genres(value: Any, overrides: Overrides) -> list[str]:
    text = clean_optional_string(value)
    if text is None:
        return []
    parts = [p.strip() for p in GENRE_SEPARATOR.split(text) if p.strip()]
    return _expand_override(parts, overrides.genre_map)


def split_countries(value: Any, overrides: Overrides) -> list[str]:
    text = clean_optional_string(value)
    if text is None:
        return []
    # Default parsing only splits commas. Explicit overrides may split a known
    # composite source value into several countries without hidden guessing.
    parts = [p.strip() for p in COUNTRY_SEPARATOR.split(text) if p.strip()]
    return _expand_override(parts, overrides.country_map)


def load_overrides(path: Path | None) -> Overrides:
    if path is None:
        return Overrides.empty()
    data = json.loads(path.read_text(encoding="utf-8"))
    return Overrides(
        genre_map=dict(data.get("genre_map", {})),
        country_map=dict(data.get("country_map", {})),
        person_map=dict(data.get("person_map", {})),
        title_map=dict(data.get("title_map", {})),
    )


def rows_as_dicts(workbook: CalamineWorkbook, sheet_name: str) -> list[dict[str, Any]]:
    rows = workbook.get_sheet_by_name(sheet_name).to_python(skip_empty_area=False)
    if not rows:
        return []
    headers = [str(x).strip() if x is not None else "" for x in rows[0]]
    out: list[dict[str, Any]] = []
    for row in rows[1:]:
        padded = list(row) + [None] * max(0, len(headers) - len(row))
        out.append(dict(zip(headers, padded, strict=False)))
    return out


def load_image_map(
    workbook: CalamineWorkbook,
    sheet_name: str,
    key_header: str,
    value_header: str,
) -> tuple[dict[str, str], dict[str, list[str]]]:
    unique: dict[str, str] = {}
    collisions: dict[str, list[str]] = defaultdict(list)
    for row in rows_as_dicts(workbook, sheet_name):
        key = clean_optional_string(row.get(key_header))
        url = clean_optional_string(row.get(value_header))
        if not key or not url:
            continue
        norm = normalize_for_match(key)
        if norm in unique and unique[norm] != url:
            collisions[norm].extend([unique[norm], url])
        else:
            unique[norm] = url
    for norm, values in collisions.items():
        unique.pop(norm, None)
        collisions[norm] = sorted(set(values))
    return unique, collisions


def parse_workbook(
    path: Path,
    overrides: Overrides,
    include_images: bool,
) -> tuple[list[ParsedMovie], list[WarningItem], dict[str, str]]:
    workbook = CalamineWorkbook.from_path(str(path))
    if MAIN_SHEET not in workbook.sheet_names:
        raise RuntimeError(f"Sheet {MAIN_SHEET!r} not found; found {workbook.sheet_names}")

    rows = rows_as_dicts(workbook, MAIN_SHEET)
    if not rows:
        raise RuntimeError("Main movie sheet is empty")

    present_headers = set(rows[0])
    missing = REQUIRED_HEADERS - present_headers
    if missing:
        raise RuntimeError(f"Missing required headers: {sorted(missing)}")

    poster_map: dict[str, str] = {}
    poster_collisions: dict[str, list[str]] = {}
    if include_images and POSTER_SHEET in workbook.sheet_names:
        poster_map, poster_collisions = load_image_map(
            workbook, POSTER_SHEET, "Film", "Foto"
        )

    warnings: list[WarningItem] = []
    movies: list[ParsedMovie] = []

    for index, row in enumerate(rows, start=2):
        excel_id = clean_optional_string(row.get("ID"))
        title = clean_optional_string(row.get("Titel"))
        if not title:
            warnings.append(
                WarningItem(index, excel_id, "Titel", row.get("Titel"), "missing_title", "Row skipped: no title", "error")
            )
            continue
        title = overrides.title_map.get(title, title)

        original_title = clean_optional_string(row.get("Originaltitel"))
        year = parse_int(row.get("Jahr"))
        runtime = parse_runtime_minutes(row.get("Playtime"))
        genres = split_genres(row.get("Genre"), overrides)
        countries = split_countries(row.get("Produktionsland"), overrides)

        director = clean_optional_string(row.get("Regisseur"))
        if director:
            director = overrides.person_map.get(director, director)

        cast: list[str] = []
        for column in ("Hauptdarsteller 1", "Hauptdarsteller 2", "Hauptdarsteller 3"):
            person = clean_optional_string(row.get(column))
            if person:
                cast.append(overrides.person_map.get(person, person))

        imdb = parse_decimal(row.get("IMDB (0-10)"))
        tomatometer = parse_rt_percent(row.get("Tomatometer"))
        popcornmeter = parse_rt_percent(row.get("Popcornmeter"))
        rating_raw, rating_0_100 = parse_star_rating(row.get("Eigene Bewertung"))

        poster_url = None
        if include_images:
            poster_url = poster_map.get(normalize_for_match(title))
            if normalize_for_match(title) in poster_collisions:
                warnings.append(
                    WarningItem(index, excel_id, "Poster", title, "poster_collision", "Multiple poster URLs match this title; poster not imported")
                )

        movie = ParsedMovie(
            row=index,
            excel_id=excel_id,
            title=title,
            original_title=original_title,
            release_year=year,
            runtime_minutes=runtime,
            overview=clean_optional_string(row.get("Kurzbeschreibung")),
            editorial_note=clean_optional_string(row.get("Anmerkung")),
            genres=genres,
            countries=countries,
            director=director,
            cast=cast,
            imdb=imdb,
            tomatometer=tomatometer,
            popcornmeter=popcornmeter,
            justwatch_url=clean_optional_string(row.get("JustWatch")),
            own_rating_raw=rating_raw,
            own_rating_0_100=rating_0_100,
            watched=parse_watched(row.get("Gesehen?")),
            poster_url=poster_url,
        )
        movies.append(movie)

        if year is None:
            warnings.append(WarningItem(index, excel_id, "Jahr", row.get("Jahr"), "invalid_year", "No valid release year"))
        elif not 1880 <= year <= 2200:
            warnings.append(WarningItem(index, excel_id, "Jahr", year, "year_out_of_range", "Release year outside schema range", "error"))

        if row.get("Playtime") not in (None, "") and runtime is None:
            warnings.append(WarningItem(index, excel_id, "Playtime", row.get("Playtime"), "invalid_runtime", "Could not convert runtime"))
        elif runtime is not None and runtime <= 0:
            warnings.append(WarningItem(index, excel_id, "Playtime", runtime, "runtime_nonpositive", "Runtime must be positive", "error"))

        if imdb is not None and not Decimal("0") <= imdb <= Decimal("10"):
            warnings.append(WarningItem(index, excel_id, "IMDB (0-10)", str(imdb), "imdb_out_of_range", "IMDb score outside 0..10", "error"))
        for field, rating in (("Tomatometer", tomatometer), ("Popcornmeter", popcornmeter)):
            if rating is not None and not Decimal("0") <= rating <= Decimal("100"):
                warnings.append(WarningItem(index, excel_id, field, str(rating), "rt_out_of_range", f"{field} outside 0..100", "error"))

        raw_genre = clean_optional_string(row.get("Genre"))
        if raw_genre:
            for genre in genres:
                if len(genre) > 100:
                    warnings.append(WarningItem(index, excel_id, "Genre", genre, "genre_too_long", "Genre exceeds schema length", "error"))
                if normalize_for_match(genre) in normalize_for_match(title):
                    warnings.append(WarningItem(index, excel_id, "Genre", genre, "genre_looks_like_title", "Genre resembles the movie title; verify manually"))

        raw_country = clean_optional_string(row.get("Produktionsland"))
        if raw_country and "." in raw_country and "," not in raw_country:
            warnings.append(WarningItem(index, excel_id, "Produktionsland", raw_country, "suspicious_country_separator", "Country value contains a period but no comma; not auto-corrected"))

        if rating_raw is not None and rating_0_100 is None:
            warnings.append(WarningItem(index, excel_id, "Eigene Bewertung", rating_raw, "unknown_legacy_rating", "Could not map legacy star rating to 0..100"))

    # Workbook-internal duplicate candidates.
    by_key: dict[tuple[str, int | None, str], list[ParsedMovie]] = defaultdict(list)
    for movie in movies:
        by_key[movie.duplicate_key].append(movie)
    for key, candidates in by_key.items():
        if len(candidates) > 1:
            ids = [m.excel_id for m in candidates]
            rows_ = [m.row for m in candidates]
            for m in candidates:
                warnings.append(
                    WarningItem(
                        m.row,
                        m.excel_id,
                        "duplicate_key",
                        {"key": key, "excel_ids": ids, "rows": rows_},
                        "duplicate_candidate",
                        "Multiple workbook rows share normalized title/year/original-title",
                        "error",
                    )
                )

    image_meta = {
        "poster_matches": str(sum(1 for m in movies if m.poster_url)),
        "poster_collision_keys": str(len(poster_collisions)),
    }
    return movies, warnings, image_meta


def write_reports(
    report_dir: Path,
    source_file: Path,
    movies: list[ParsedMovie],
    warnings: list[WarningItem],
    image_meta: dict[str, str],
) -> None:
    report_dir.mkdir(parents=True, exist_ok=True)

    warning_counts = Counter(w.code for w in warnings)
    severity_counts = Counter(w.severity for w in warnings)
    genres = sorted({g for m in movies for g in m.genres}, key=str.casefold)
    countries = sorted({c for m in movies for c in m.countries}, key=str.casefold)
    people = sorted(
        {p for m in movies for p in ([m.director] if m.director else []) + m.cast},
        key=str.casefold,
    )

    summary = {
        "source_file": str(source_file),
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "movie_rows_parsed": len(movies),
        "warning_count": len(warnings),
        "warnings_by_severity": dict(severity_counts),
        "warnings_by_code": dict(sorted(warning_counts.items())),
        "unique_genres": len(genres),
        "unique_countries": len(countries),
        "unique_people_in_director_or_top3_cast": len(people),
        "pending_personal_rows": sum(1 for m in movies if m.own_rating_raw is not None or m.watched),
        "image_enrichment": image_meta,
    }
    (report_dir / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    with (report_dir / "warnings.csv").open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["severity", "code", "row", "excel_id", "field", "value", "message"])
        for w in warnings:
            writer.writerow([
                w.severity,
                w.code,
                w.row,
                w.excel_id,
                w.field,
                json.dumps(w.value, ensure_ascii=False) if isinstance(w.value, (dict, list)) else w.value,
                w.message,
            ])

    with (report_dir / "movies_preview.csv").open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow([
            "excel_id", "row", "title", "original_title", "release_year",
            "runtime_minutes", "genres", "countries", "director", "cast",
            "imdb", "tomatometer", "popcornmeter", "justwatch_url", "poster_url",
        ])
        for m in movies:
            writer.writerow([
                m.excel_id,
                m.row,
                m.title,
                m.original_title,
                m.release_year,
                m.runtime_minutes,
                " | ".join(m.genres),
                " | ".join(m.countries),
                m.director,
                " | ".join(m.cast),
                m.imdb,
                m.tomatometer,
                m.popcornmeter,
                m.justwatch_url,
                m.poster_url,
            ])

    with (report_dir / "personal_pending.csv").open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["excel_id", "title", "release_year", "legacy_rating", "rating_0_100", "watched"])
        for m in movies:
            if m.own_rating_raw is not None or m.watched:
                writer.writerow([
                    m.excel_id,
                    m.title,
                    m.release_year,
                    m.own_rating_raw,
                    m.own_rating_0_100,
                    m.watched,
                ])

    for filename, values in (
        ("genres.txt", genres),
        ("countries.txt", countries),
        ("people.txt", people),
    ):
        (report_dir / filename).write_text("\n".join(values) + "\n", encoding="utf-8")


async def commit_movies(
    movies: list[ParsedMovie],
    *,
    skip_duplicates: bool,
) -> dict[str, int]:
    """Write parsed public movie data using the agreed async SQLAlchemy models."""
    try:
        from sqlalchemy import func, select
        from sqlalchemy.orm import selectinload

        from app.db.session import async_session_factory
        from app.models import (
            AuditLog,
            Country,
            ExternalRating,
            Genre,
            Movie,
            MovieCountry,
            MovieGenre,
            MovieLink,
            MoviePerson,
            Person,
        )
        from app.services.movie_loader import movie_full_options
        from app.services.movie_snapshot import movie_snapshot
    except ImportError as exc:
        raise RuntimeError(
            "--commit requires this script to run from the backend project where "
            "app.db.session and app.models are importable."
        ) from exc

    stats = Counter()
    now = datetime.now(timezone.utc)
    created_movie_ids = []

    async with async_session_factory() as session:
        # Preload small lookup tables once; name matching here is case-insensitive.
        genre_rows = (await session.execute(select(Genre))).scalars().all()
        country_rows = (await session.execute(select(Country))).scalars().all()
        person_rows = (await session.execute(select(Person))).scalars().all()

        genre_cache = {normalize_for_match(g.name): g for g in genre_rows}
        country_cache = {normalize_for_match(c.name): c for c in country_rows}
        person_cache = {normalize_for_match(p.name): p for p in person_rows}

        for data in movies:
            # Conservative DB duplicate check: same case-insensitive title + release year.
            stmt = select(Movie).where(
                func.lower(Movie.title) == data.title.casefold(),
                Movie.release_year == data.release_year,
            )
            existing = (await session.execute(stmt)).scalars().first()
            if existing is not None:
                stats["db_duplicate_candidates"] += 1
                if skip_duplicates:
                    stats["skipped_duplicates"] += 1
                    continue
                raise RuntimeError(
                    f"DB duplicate candidate for {data.title!r} ({data.release_year}); "
                    "rerun dry-run / inspect data or use --skip-duplicates explicitly."
                )

            movie = Movie(
                title=data.title,
                original_title=data.original_title,
                release_year=data.release_year,
                release_date=None,
                runtime_minutes=data.runtime_minutes,
                overview=data.overview,
                editorial_note=data.editorial_note,
                poster_url=data.poster_url,
                poster_path=None,
                tmdb_id=None,
                imdb_id=None,
                created_by_id=None,
            )
            session.add(movie)
            await session.flush()

            seen_genres: set[str] = set()
            for genre_name in data.genres:
                key = normalize_for_match(genre_name)

                if key in seen_genres:
                    continue
                seen_genres.add(key)

                genre = genre_cache.get(key)
                if genre is None:
                    genre = Genre(name=genre_name)
                    session.add(genre)
                    await session.flush()
                    genre_cache[key] = genre
                    stats["genres_created"] += 1

                session.add(MovieGenre(movie_id=movie.id, genre_id=genre.id))

            for country_name in data.countries:
                key = normalize_for_match(country_name)
                country = country_cache.get(key)
                if country is None:
                    country = Country(name=country_name, iso2=None, iso3=None)
                    session.add(country)
                    await session.flush()
                    country_cache[key] = country
                    stats["countries_created"] += 1
                session.add(MovieCountry(movie_id=movie.id, country_id=country.id))

            if data.director:
                key = normalize_for_match(data.director)
                person = person_cache.get(key)
                if person is None:
                    person = Person(name=data.director)
                    session.add(person)
                    await session.flush()
                    person_cache[key] = person
                    stats["people_created"] += 1
                session.add(
                    MoviePerson(
                        movie_id=movie.id,
                        person_id=person.id,
                        credit_type="crew",
                        job="Director",
                    )
                )

            for order, actor_name in enumerate(data.cast):
                key = normalize_for_match(actor_name)
                person = person_cache.get(key)
                if person is None:
                    person = Person(name=actor_name)
                    session.add(person)
                    await session.flush()
                    person_cache[key] = person
                    stats["people_created"] += 1
                session.add(
                    MoviePerson(
                        movie_id=movie.id,
                        person_id=person.id,
                        credit_type="cast",
                        character_name=None,
                        billing_order=order,
                    )
                )

            if data.imdb is not None:
                session.add(
                    ExternalRating(
                        movie_id=movie.id,
                        source="imdb",
                        rating_type="audience",
                        value=data.imdb,
                        scale=Decimal("10"),
                        vote_count=None,
                        source_url=None,
                        retrieved_at=now,
                    )
                )
            if data.tomatometer is not None:
                session.add(
                    ExternalRating(
                        movie_id=movie.id,
                        source="rotten_tomatoes",
                        rating_type="critics",
                        value=data.tomatometer,
                        scale=Decimal("100"),
                        vote_count=None,
                        source_url=None,
                        retrieved_at=now,
                    )
                )
            if data.popcornmeter is not None:
                session.add(
                    ExternalRating(
                        movie_id=movie.id,
                        source="rotten_tomatoes",
                        rating_type="audience",
                        value=data.popcornmeter,
                        scale=Decimal("100"),
                        vote_count=None,
                        source_url=None,
                        retrieved_at=now,
                    )
                )

            if data.justwatch_url:
                session.add(
                    MovieLink(
                        movie_id=movie.id,
                        source="justwatch",
                        url=data.justwatch_url,
                        label="JustWatch",
                    )
                )

            created_movie_ids.append(movie.id)
            stats["movies_created"] += 1

        # Seed import is part of the audit history. Batch-load the complete
        # created Movie graphs after all relationship rows have been flushed,
        # then append one system CREATE audit row per imported film.
        await session.flush()
        if created_movie_ids:
            stmt = (
                select(Movie)
                .where(Movie.id.in_(created_movie_ids))
                .options(*movie_full_options())
            )
            imported_movies = (await session.execute(stmt)).scalars().unique().all()
            for imported_movie in imported_movies:
                session.add(
                    AuditLog(
                        actor_user_id=None,
                        action="CREATE",
                        entity_type="movie",
                        entity_id=imported_movie.id,
                        before_data=None,
                        after_data=movie_snapshot(imported_movie),
                    )
                )
            stats["audit_logs_created"] += len(imported_movies)

        await session.commit()

    return dict(stats)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Dry-run/commit importer for FILME-V1.1.xlsm",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument("xlsx", type=Path, help="Path to FILME-V1.1.xlsm")
    parser.add_argument(
        "--report-dir",
        type=Path,
        default=Path("import_report"),
        help="Directory for dry-run reports",
    )
    parser.add_argument(
        "--overrides",
        type=Path,
        default=None,
        help="Optional JSON with explicit genre/country/person/title replacements",
    )
    parser.add_argument(
        "--include-images",
        action="store_true",
        help="Also map poster URLs from 'Fotos FP'. Off by default.",
    )
    parser.add_argument(
        "--commit",
        action="store_true",
        help="Actually write public movie data to PostgreSQL. Default is dry-run only.",
    )
    parser.add_argument(
        "--allow-warnings",
        action="store_true",
        help="Permit commit when non-error warnings remain. Errors still block commit.",
    )
    parser.add_argument(
        "--skip-duplicates",
        action="store_true",
        help="On commit, skip DB duplicate candidates instead of aborting.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if not args.xlsx.exists():
        print(f"ERROR: file not found: {args.xlsx}", file=sys.stderr)
        return 2

    overrides = load_overrides(args.overrides)
    movies, warnings, image_meta = parse_workbook(
        args.xlsx,
        overrides=overrides,
        include_images=args.include_images,
    )
    write_reports(args.report_dir, args.xlsx, movies, warnings, image_meta)

    errors = [w for w in warnings if w.severity == "error"]
    ordinary_warnings = [w for w in warnings if w.severity != "error"]

    print(f"Parsed movies: {len(movies)}")
    print(f"Warnings:      {len(ordinary_warnings)}")
    print(f"Errors:        {len(errors)}")
    print(f"Reports:       {args.report_dir.resolve()}")
    print(f"Mode:          {'COMMIT' if args.commit else 'DRY RUN'}")

    if not args.commit:
        print("No database writes performed. Review summary.json and warnings.csv first.")
        return 1 if errors else 0

    if errors:
        print("COMMIT BLOCKED: error-severity findings remain in the dry-run report.", file=sys.stderr)
        return 3
    if ordinary_warnings and not args.allow_warnings:
        print(
            "COMMIT BLOCKED: warnings remain. Review them and either provide explicit "
            "overrides or rerun with --allow-warnings if you intentionally accept them.",
            file=sys.stderr,
        )
        return 4

    stats = asyncio.run(
        commit_movies(
            movies,
            skip_duplicates=args.skip_duplicates,
        )
    )
    print("Commit complete:")
    for key, value in sorted(stats.items()):
        print(f"  {key}: {value}")
    print("Personal ratings/watched state were NOT imported; see personal_pending.csv.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
